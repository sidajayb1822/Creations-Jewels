import os
from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel,
    QPushButton, QFileDialog, QMessageBox,
    QGroupBox, QGridLayout, QTabWidget, QLineEdit,
)
from PySide6.QtCore import QThread, Signal, QObject
from PySide6.QtGui import QFont, QColor

from app.ui.bom_compare.drop_zone import DropZone
from app.ui.bom_compare.comparison_table import ComparisonTable
from app.ui.bom_compare.line_detail import LineDetailDialog
from app.core.excel_parser import parse_workbook
from app.core.comparator import compare, summary_stats, ComparisonRow
from app.core.db import db, load_config, DEFAULT_METAL_LOSS_PCT


class _Worker(QObject):
    """Run parse + compare for every design in a workbook, off the UI thread."""
    done = Signal(list, str)   # [ {label, rows, stats}, ... ], order_info
    error = Signal(str)

    def __init__(self, file_path: str):
        super().__init__()
        self._path = file_path

    def run(self):
        try:
            docs = parse_workbook(self._path)

            # Company code + costing settings come from config / the first design.
            cfg = load_config()
            company_code = (cfg.get("company_code") or "").strip()
            base_company_code = (cfg.get("base_company_code") or "").strip()
            customer = docs[0].header.customer if docs else ""
            if not company_code and customer and db.is_connected():
                company_code = db.find_company_code(customer)

            # Header shows the loss for the first metal's category; the actual
            # loss is applied per metal category inside compare().
            first_cat = next((d.metals[0].category for d in docs if d.metals), "")
            loss_pct = db.get_metal_loss_pct(company_code, first_cat)
            loss_source = "Emperor"
            if loss_pct is None:
                loss_pct = float(cfg.get("metal_loss_pct", DEFAULT_METAL_LOSS_PCT))
                loss_source = "default"
            multiplier = db.get_customer_multiplier(company_code)

            results = []
            for d in docs:
                rows = compare(d, db, company_code=company_code, loss_pct=loss_pct,
                               base_company_code=base_company_code)
                stats = summary_stats(rows, multiplier=multiplier)
                results.append({"label": d.design_label, "sheet": d.sheet_name,
                                "rows": rows, "stats": stats})

            order_no = docs[0].header.order_no if docs else ""
            order_info = (
                f"{order_no}  |  {customer}  |  {len(docs)} design(s)"
                + (f"  |  Co: {company_code}" if company_code else "")
                + (f"  |  Base: {base_company_code}"
                   if base_company_code and base_company_code != company_code else "")
                + f"  |  Loss: {loss_pct:g}% ({loss_source})"
                + (f"  |  ×{multiplier:g}" if multiplier != 1.0 else "")
            )
            self.done.emit(results, order_info)
        except Exception as e:
            self.error.emit(str(e))


class _DesignComparePanel(QWidget):
    """One design's comparison table + summary bar. Reused per sub-tab."""

    def __init__(self, on_edit_formula=None, parent=None):
        super().__init__(parent)
        self._rows: list[ComparisonRow] = []
        self._on_edit_formula = on_edit_formula
        self.sheet_name: str = ""

        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 8, 0, 0)
        layout.setSpacing(8)

        self.table = ComparisonTable()
        self.table.cellDoubleClicked.connect(self._open_line_detail)
        layout.addWidget(self.table, 1)

        layout.addWidget(self._build_summary_bar())

        hint = QLabel("Double-click a line to see how it is calculated and edit its formula."
                      "    * = rate taken from the base chart (customer has no own entry).")
        hint.setStyleSheet("color: #94a3b8; font-size: 11px;")
        layout.addWidget(hint)

    def _build_summary_bar(self) -> QGroupBox:
        group = QGroupBox("Summary")
        grid = QGridLayout(group)
        grid.setSpacing(14)

        def _stat(label: str, col: int) -> QLabel:
            lbl = QLabel(label)
            lbl.setStyleSheet("color: #64748b; font-size: 11px;")
            val = QLabel("—")
            val.setFont(QFont("Segoe UI", 12, QFont.Weight.Bold))
            grid.addWidget(lbl, 0, col)
            grid.addWidget(val, 1, col)
            return val

        self._sum_template = _stat("Template Total ($)", 0)
        self._sum_master = _stat("Master Total ($)", 1)
        self._sum_diff = _stat("Difference ($)", 2)
        self._sum_pct = _stat("Difference (%)", 3)
        self._sum_ok = _stat("OK", 4)
        self._sum_warn = _stat("Warning", 5)
        self._sum_err = _stat("Over Threshold", 6)
        self._sum_miss = _stat("Missing", 7)
        return group

    def load(self, rows: list[ComparisonRow], stats: dict):
        self._rows = rows
        self.table.load(rows)
        c = stats["counts"]
        self._sum_template.setText(f"$ {stats['total_template']:,.2f}")
        self._sum_master.setText(f"$ {stats['total_master']:,.2f}")
        diff, pct = stats["total_diff"], stats["total_pct"]
        color = "#10b981" if abs(pct) <= 5 else ("#d97706" if abs(pct) <= 15 else "#dc2626")
        self._sum_diff.setText(f"$ {diff:+,.2f}")
        self._sum_diff.setStyleSheet(f"color: {color}; font-size: 12px; font-weight: bold;")
        self._sum_pct.setText(f"{pct:+.1f}%")
        self._sum_pct.setStyleSheet(f"color: {color}; font-size: 12px; font-weight: bold;")
        self._sum_ok.setText(str(c.get("match", 0)))
        self._sum_warn.setText(str(c.get("minor", 0)))
        self._sum_err.setText(str(c.get("major", 0)))
        self._sum_miss.setText(str(c.get("missing", 0)))

    def rows(self) -> list[ComparisonRow]:
        return self._rows

    def _open_line_detail(self, table_row: int, _col: int):
        row = self.table.row_at(table_row)
        if row is None:
            return
        dlg = LineDetailDialog(row, self)
        dlg.exec()
        if dlg.changed and self._on_edit_formula:
            self._on_edit_formula()


class BOMCompareWidget(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        self._thread: QThread | None = None
        self._last_path: str = ""
        self._build_ui()

    def _build_ui(self):
        layout = QVBoxLayout(self)
        layout.setSpacing(10)
        layout.setContentsMargins(16, 16, 16, 16)

        drop_row = QHBoxLayout()
        self.drop_zone = DropZone()
        self.drop_zone.file_dropped.connect(self._on_file)
        drop_row.addWidget(self.drop_zone)

        # Browse + Refresh stacked vertically so neither gets squished.
        btn_col = QVBoxLayout()
        btn_col.setSpacing(6)

        browse_btn = QPushButton("Browse…")
        browse_btn.setObjectName("outlineBtn")
        browse_btn.setFixedWidth(100)
        browse_btn.clicked.connect(self._browse)
        btn_col.addWidget(browse_btn)

        self.refresh_btn = QPushButton("↻ Refresh")
        self.refresh_btn.setObjectName("outlineBtn")
        self.refresh_btn.setFixedWidth(100)
        self.refresh_btn.setToolTip("Re-pull rates from Emperor and re-run the "
                                    "comparison for the current file.")
        self.refresh_btn.setEnabled(False)
        self.refresh_btn.clicked.connect(self._refresh)
        btn_col.addWidget(self.refresh_btn)

        drop_row.addLayout(btn_col)
        layout.addLayout(drop_row)

        self.order_label = QLabel("No file loaded")
        self.order_label.setStyleSheet("color: #64748b; font-style: italic;")
        layout.addWidget(self.order_label)

        # Filter the design tabs by name (shown only for multi-design files).
        self.search_edit = QLineEdit()
        self.search_edit.setPlaceholderText("Search designs…")
        self.search_edit.setClearButtonEnabled(True)
        self.search_edit.textChanged.connect(self._filter_tabs)
        self.search_edit.hide()
        layout.addWidget(self.search_edit)

        # One sub-tab per design.
        self.tabs = QTabWidget()
        self.tabs.setObjectName("designTabs")
        layout.addWidget(self.tabs, 1)

        btn_row = QHBoxLayout()
        self.export_btn = QPushButton("Export Highlighted Excel")
        self.export_btn.setEnabled(False)
        self.export_btn.clicked.connect(self._export)
        btn_row.addStretch()
        btn_row.addWidget(self.export_btn)
        layout.addLayout(btn_row)

    def _browse(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Open Emperor BOM File", "", "Excel Files (*.xlsx)"
        )
        if path:
            self._on_file(path)

    def _on_file(self, path: str):
        self._last_path = path
        self.order_label.setText(f"Loading  {os.path.basename(path)} …")
        self.export_btn.setEnabled(False)
        self.tabs.clear()

        self._thread = QThread()
        self._worker = _Worker(path)
        self._worker.moveToThread(self._thread)
        self._thread.started.connect(self._worker.run)
        self._worker.done.connect(self._on_done)
        self._worker.error.connect(self._on_error)
        self._worker.done.connect(self._thread.quit)
        self._worker.error.connect(self._thread.quit)
        self._thread.start()

    def _on_done(self, results: list, order_info: str):
        self.tabs.clear()
        diff_designs = 0
        for res in results:
            panel = _DesignComparePanel(on_edit_formula=self._reload)
            panel.sheet_name = res.get("sheet", "")
            panel.load(res["rows"], res["stats"])
            idx = self.tabs.addTab(panel, res["label"])

            # Highlight tabs whose comparison has a real price difference.
            c = res["stats"].get("counts", {})
            major, minor, missing = (c.get("major", 0), c.get("minor", 0),
                                     c.get("missing", 0))
            if major:
                self.tabs.tabBar().setTabTextColor(idx, QColor("#dc2626"))
            elif minor:
                self.tabs.tabBar().setTabTextColor(idx, QColor("#d97706"))
            if major or minor:
                diff_designs += 1
            parts = []
            if major:
                parts.append(f"{major} over threshold")
            if minor:
                parts.append(f"{minor} minor")
            if missing:
                parts.append(f"{missing} missing")
            self.tabs.setTabToolTip(idx, ", ".join(parts) if parts else "all match")

        if diff_designs:
            order_info += f"  |  {diff_designs} design(s) with differences"
        self.order_label.setText(order_info)

        # Search bar only helps when there's more than one design.
        self.search_edit.blockSignals(True)
        self.search_edit.clear()
        self.search_edit.blockSignals(False)
        self.search_edit.setVisible(len(results) > 1)
        self.export_btn.setEnabled(bool(results))
        self.refresh_btn.setEnabled(bool(self._last_path))

    def _refresh(self):
        """Drop cached rates and re-run the comparison against fresh Emperor data."""
        if not self._last_path:
            return
        db.clear_caches()
        self._reload()

    def _filter_tabs(self, text: str):
        """Show only design tabs whose name contains the search text."""
        q = (text or "").strip().lower()
        first_visible = -1
        for i in range(self.tabs.count()):
            visible = q in self.tabs.tabText(i).lower()
            self.tabs.setTabVisible(i, visible)
            if visible and first_visible < 0:
                first_visible = i
        cur = self.tabs.currentIndex()
        if cur >= 0 and not self.tabs.isTabVisible(cur) and first_visible >= 0:
            self.tabs.setCurrentIndex(first_visible)

    def _on_error(self, msg: str):
        self.order_label.setText("Error loading file")
        QMessageBox.critical(self, "Parse Error", msg)

    def _reload(self):
        """Re-run the comparison after a formula edit so numbers update."""
        if self._last_path:
            self._on_file(self._last_path)

    def _export(self):
        # Export = a copy of the uploaded file with each design's rows filled by
        # their status colour; everything else in the workbook is left untouched.
        if not self._last_path:
            return
        designs: list[tuple[str, list[ComparisonRow]]] = []
        for i in range(self.tabs.count()):
            panel = self.tabs.widget(i)
            if isinstance(panel, _DesignComparePanel) and panel.sheet_name and panel.rows():
                designs.append((panel.sheet_name, panel.rows()))
        if not designs:
            return
        stem = os.path.splitext(os.path.basename(self._last_path))[0]
        path, _ = QFileDialog.getSaveFileName(
            self, "Save Highlighted Excel", f"{stem}-highlighted.xlsx",
            "Excel Files (*.xlsx)"
        )
        if not path:
            return
        try:
            n = _export_highlighted(self._last_path, designs, path)
            QMessageBox.information(
                self, "Exported",
                f"Highlighted {n} design sheet(s) — original layout preserved:\n{path}")
        except Exception as e:
            QMessageBox.critical(self, "Export Error", str(e))


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

_STATUS_FILLS = {
    "match":   "FFd4edda",
    "minor":   "FFfff3cd",
    "major":   "FFf8d7da",
    "missing": "FFe2e3e5",
}

# When several comparison lines map to one Excel row (a stone + its setting),
# the row takes the most severe status.
_STATUS_SEVERITY = {"match": 0, "missing": 1, "minor": 2, "major": 3}


def _export_highlighted(src_path: str,
                        designs: list[tuple[str, list[ComparisonRow]]],
                        out_path: str) -> int:
    """
    Save a copy of the uploaded workbook (`src_path`) to `out_path`, unchanged
    except that each design sheet's data rows are filled with their status
    colour. `designs` is [(sheet_name, rows), ...]. Returns the number of sheets
    coloured. Non-design sheets and header rows are left exactly as they were.
    """
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill

    wb = load_workbook(src_path)   # no data_only → values/formulas/styles preserved
    coloured = 0
    for sheet_name, rows in designs:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        # Worst status per physical Excel row.
        worst: dict[int, str] = {}
        for row in rows:
            sr = getattr(row, "source_row", 0)
            if not sr:
                continue
            cur = worst.get(sr)
            if cur is None or _STATUS_SEVERITY.get(row.status, 0) > _STATUS_SEVERITY.get(cur, 0):
                worst[sr] = row.status
        max_col = ws.max_column
        for row_idx, status in worst.items():
            fill = PatternFill("solid", fgColor=_STATUS_FILLS.get(status, "FFFFFFFF"))
            for col in range(1, max_col + 1):
                ws.cell(row=row_idx, column=col).fill = fill
        coloured += 1
    wb.save(out_path)
    return coloured
